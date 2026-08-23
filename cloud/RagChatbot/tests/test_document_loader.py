import io
import os
import tempfile
import pytest
from unittest.mock import MagicMock, patch

from RagChatbot.services.document_loader import (
    extract_text_from_file,
    is_supported_document_extension,
    SUPPORTED_EXTENSIONS,
)


def test_supported_extensions():
    assert is_supported_document_extension(".pdf")
    assert is_supported_document_extension("pdf")
    assert is_supported_document_extension(".docx")
    assert is_supported_document_extension(".doc")
    assert is_supported_document_extension(".txt")
    assert is_supported_document_extension(".md")
    assert is_supported_document_extension(".csv")
    assert not is_supported_document_extension(".png")
    assert not is_supported_document_extension(".exe")


def test_extract_from_plain_text():
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
        f.write("Campus Library Hours:\nMonday to Friday: 8:00 AM - 10:00 PM\nSaturday: 9:00 AM - 5:00 PM")
        tmp_path = f.name

    try:
        text = extract_text_from_file(tmp_path)
        assert "Campus Library Hours" in text
        assert "8:00 AM" in text
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_extract_from_markdown():
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as f:
        f.write("# Computer Science Faculty\n\nDean: Dr. Jane Doe\n\n## Degree Programmes\n- Bachelor of Computer Science\n- Diploma in IT")
        tmp_path = f.name

    try:
        text = extract_text_from_file(tmp_path)
        assert "# Computer Science Faculty" in text
        assert "Bachelor of Computer Science" in text
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_extract_from_csv():
    with tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, encoding="utf-8") as f:
        f.write("Course Code,Course Name,Credit Hours\nBCS101,Introduction to AI,3\nBCS102,Data Structures,4\n")
        tmp_path = f.name

    try:
        text = extract_text_from_file(tmp_path)
        assert "| Course Code | Course Name | Credit Hours |" in text
        assert "| BCS101 | Introduction to AI | 3 |" in text
        assert "| BCS102 | Data Structures | 4 |" in text
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_extract_from_pdf():
    import pypdf
    from pypdf import PdfWriter

    writer = PdfWriter()
    # Add a page with blank text or draw text
    writer.add_blank_page(width=300, height=300)
    
    with tempfile.NamedTemporaryFile("wb", suffix=".pdf", delete=False) as f:
        writer.write(f)
        tmp_path = f.name

    try:
        # Since pypdf blank page has no text, mock extract_text
        with patch("pypdf.PdfReader") as mock_reader_cls:
            mock_reader = MagicMock()
            mock_page = MagicMock()
            mock_page.extract_text.return_value = "Quest International University Exam Regulations. All candidates must bring student ID card."
            mock_reader.pages = [mock_page]
            mock_reader_cls.return_value = mock_reader

            text = extract_text_from_file(tmp_path)
            assert "## Page 1" in text
            assert "Quest International University Exam Regulations" in text
            assert "student ID card" in text
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_extract_from_docx():
    import docx

    doc = docx.Document()
    doc.add_heading("Campus Parking Guidelines", level=1)
    doc.add_paragraph("All vehicles must display valid parking permits.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Zone"
    table.cell(0, 1).text = "Fee"
    table.cell(1, 0).text = "Zone A (Students)"
    table.cell(1, 1).text = "RM 50 / semester"

    with tempfile.NamedTemporaryFile("wb", suffix=".docx", delete=False) as f:
        doc.save(f)
        tmp_path = f.name

    try:
        text = extract_text_from_file(tmp_path)
        assert "# Campus Parking Guidelines" in text
        assert "All vehicles must display valid parking permits." in text
        assert "| Zone | Fee |" in text
        assert "| Zone A (Students) | RM 50 / semester |" in text
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_extract_unsupported_file():
    with tempfile.NamedTemporaryFile("w", suffix=".xyz", delete=False) as f:
        f.write("Some invalid content")
        tmp_path = f.name

    try:
        with pytest.raises(ValueError, match="Unsupported file type"):
            extract_text_from_file(tmp_path)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
